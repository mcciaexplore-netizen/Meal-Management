import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime, timedelta
from uuid import UUID, uuid4

from .auth import audit, require_actor
from .database import retry_transaction
from .email_queue import RESEND_LIMIT
from .errors import DomainError
from .paging import newest_first, next_cursor, parse_cursor
from .security import normalize_email, normalize_phone, utc_naive

HEADERS = ("Sr. No.", "Person Name", "Company Name", "Phone", "Email")
MAX_ROWS = 200
MAX_FILE_BYTES = 1_048_576
MAX_UNZIPPED_BYTES = 20 * 1_048_576
MAX_ZIP_ENTRIES = 200
CSV_TYPES = {"text/csv", "application/csv", "text/plain", "application/vnd.ms-excel"}
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MESSAGES = {
    "REQUIRED": "is required",
    "TOO_LONG": "is longer than 150 characters",
    "INVALID_SR_NO": "must be a whole number from 1",
    "DUPLICATE_SR_NO": "is used by another row",
    "INVALID_PHONE": "is not a valid phone number",
    "INVALID_EMAIL": "is not a valid email address",
    "DUPLICATE_PHONE": "is also used by another row",
    "DUPLICATE_EMAIL": "is also used by another row",
    "INVALID_TEXT": "contains characters that are not allowed",
}


def _header_key(value):
    return " ".join(str(value or "").replace("﻿", "").split()).casefold()


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _rows_from_csv(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DomainError("VISITOR_FILE_ENCODING") from None
    return [[_cell(value) for value in row] for row in csv.reader(io.StringIO(text))]


def _rows_from_xlsx(data):
    from openpyxl import load_workbook

    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
        entries = archive.infolist()
    except zipfile.BadZipFile:
        raise DomainError("VISITOR_FILE_INVALID") from None
    names = {entry.filename for entry in entries}
    if (len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_UNZIPPED_BYTES
            or "xl/workbook.xml" not in names):
        raise DomainError("VISITOR_FILE_INVALID")
    if any(name.startswith("xl/externalLinks/") for name in names):
        raise DomainError("VISITOR_FILE_EXTERNAL_LINKS")
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        raise DomainError("VISITOR_FILE_MACROS")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_vba=False, keep_links=False)
    except Exception:
        raise DomainError("VISITOR_FILE_INVALID") from None
    try:
        if getattr(workbook, "_external_links", None):
            raise DomainError("VISITOR_FILE_EXTERNAL_LINKS")
        rows = []
        for row in workbook.worksheets[0].iter_rows():
            values = []
            for cell in row:
                if getattr(cell, "data_type", None) == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    raise DomainError("VISITOR_FILE_FORMULAS")
                values.append(_cell(cell.value))
            rows.append(values)
            if len(rows) > MAX_ROWS + 50:
                raise DomainError("VISITOR_FILE_TOO_MANY_ROWS")
        return rows
    finally:
        workbook.close()


def read_rows(data, content_type, filename):
    if not isinstance(data, (bytes, bytearray)) or not data:
        raise DomainError("VISITOR_FILE_REQUIRED")
    if len(data) > MAX_FILE_BYTES:
        raise DomainError("VISITOR_FILE_TOO_LARGE")
    name = (filename or "").strip().lower()
    kind = (content_type or "").split(";")[0].strip().lower()
    if name.endswith(".xlsx") and kind in {XLSX_TYPE, "application/octet-stream"} and bytes(data[:4]) == b"PK\x03\x04":
        table = _rows_from_xlsx(bytes(data))
    elif name.endswith(".csv") and kind in CSV_TYPES | {"application/octet-stream"}:
        table = _rows_from_csv(bytes(data))
    else:
        raise DomainError("VISITOR_FILE_TYPE")
    while table and not any(table[-1]):
        table.pop()
    if not table:
        raise DomainError("VISITOR_FILE_EMPTY")
    header = [value for value in table[0]]
    while header and not header[-1]:
        header.pop()
    if [_header_key(value) for value in header] != [_header_key(value) for value in HEADERS]:
        raise DomainError("VISITOR_FILE_HEADERS")
    rows = []
    for number, values in enumerate(table[1:], start=2):
        if not any(values):
            continue
        if any(values[len(HEADERS):]):
            raise DomainError("VISITOR_FILE_HEADERS")
        values = (values + [""] * len(HEADERS))[:len(HEADERS)]
        rows.append({"row": number, "sr_no": values[0], "person_name": values[1], "company_name": values[2],
                     "phone": values[3], "email": values[4]})
    if not rows:
        raise DomainError("VISITOR_FILE_EMPTY")
    if len(rows) > MAX_ROWS:
        raise DomainError("VISITOR_FILE_TOO_MANY_ROWS")
    return rows


def validate_rows(rows):
    if not isinstance(rows, list) or not 1 <= len(rows) <= MAX_ROWS:
        raise DomainError("VISITOR_FILE_TOO_MANY_ROWS" if isinstance(rows, list) and rows else "VISITOR_FILE_EMPTY")
    clean, errors = [], []
    seen = {"sr_no": {}, "phone": {}, "email": {}}

    def error(row, column, code):
        errors.append({"row": row, "column": column, "code": code, "message": f"{column} {MESSAGES[code]}."})

    for index, raw in enumerate(rows):
        if not isinstance(raw, dict):
            raise DomainError("INVALID_VISITOR_ROWS")
        number = raw.get("row") if type(raw.get("row")) is int and 1 <= raw.get("row") <= 100_000 else index + 2
        values = {key: raw.get(key) for key in ("sr_no", "person_name", "company_name", "phone", "email")}
        if any(value is not None and not isinstance(value, (str, int)) or isinstance(value, bool) for value in values.values()):
            raise DomainError("INVALID_VISITOR_ROWS")
        text = {key: "" if value is None else str(value).strip() for key, value in values.items()}
        row = {"row": number}
        if not text["sr_no"].isdigit() or int(text["sr_no"]) < 1 or len(text["sr_no"]) > 6:
            error(number, "Sr. No.", "INVALID_SR_NO")
        else:
            row["sr_no"] = int(text["sr_no"])
        for key, column in (("person_name", "Person Name"), ("company_name", "Company Name")):
            if not text[key]:
                error(number, column, "REQUIRED")
            elif len(text[key]) > 150:
                error(number, column, "TOO_LONG")
            elif any(ord(character) < 32 or ord(character) == 127 for character in text[key]):
                error(number, column, "INVALID_TEXT")
            else:
                row[key] = text[key]
        if not text["phone"]:
            error(number, "Phone", "REQUIRED")
        else:
            try:
                row["phone"] = normalize_phone(text["phone"], "VISITOR_PHONE")
            except DomainError:
                error(number, "Phone", "INVALID_PHONE")
        row["email"] = None
        if text["email"]:
            try:
                row["email"] = normalize_email(text["email"])
            except DomainError:
                error(number, "Email", "INVALID_EMAIL")
        for key, column, code in (("sr_no", "Sr. No.", "DUPLICATE_SR_NO"), ("phone", "Phone", "DUPLICATE_PHONE"),
                                  ("email", "Email", "DUPLICATE_EMAIL")):
            value = row.get(key)
            if value is None:
                continue
            if key == "phone":
                identity = "".join(character for character in value if character.isdigit())
            else:
                identity = value.casefold() if isinstance(value, str) else value
            if identity in seen[key]:
                first = seen[key][identity]
                if first is not None:
                    error(first, column, code)
                    seen[key][identity] = None
                error(number, column, code)
            else:
                seen[key][identity] = number
        clean.append(row)
    errors.sort(key=lambda item: (item["row"], item["column"]))
    return clean, errors


def _record(row):
    result = dict(row)
    for key, value in result.items():
        if isinstance(value, datetime) and value.tzinfo is None:
            result[key] = value.isoformat() + "Z"
    return result


class VisitorBatchService:
    def __init__(self, db, qr, email_daily_limit=450):
        self.db = db
        self.qr = qr
        self.daily_limit = email_daily_limit

    def preview(self, context, data, content_type, filename):
        rows = read_rows(data, content_type, filename)
        with self.db.transaction() as tx:
            require_actor(tx, context, {"ADMIN"})
        clean, errors = validate_rows(rows)
        return {"rows": rows, "errors": errors, "row_count": len(rows), "valid": not errors,
                "email_count": sum(1 for row in clean if row.get("email"))}

    def _settings(self, expires_at, meal_type_ids, department_id, event_name, source_name):
        if not isinstance(expires_at, datetime):
            raise DomainError("VISITOR_BATCH_EXPIRY_REQUIRED")
        expiry = utc_naive(expires_at)
        if (not isinstance(meal_type_ids, list) or not 1 <= len(meal_type_ids) <= 50
                or any(type(value) is not int or value < 1 for value in meal_type_ids)
                or len(set(meal_type_ids)) != len(meal_type_ids)):
            raise DomainError("INVALID_MEAL_TYPE_IDS")
        if type(department_id) is not int or not 1 <= department_id <= 2**64 - 1:
            raise DomainError("INVALID_DEPARTMENT_ID")
        if event_name is not None and (not isinstance(event_name, str) or len(event_name.strip()) > 150):
            raise DomainError("INVALID_EVENT_NAME")
        if not isinstance(source_name, str) or not source_name.strip() or len(source_name) > 255:
            raise DomainError("INVALID_SOURCE_NAME")
        return {"expires_at": expiry, "meal_type_ids": sorted(meal_type_ids), "department_id": department_id,
                "event_name": (event_name or "").strip() or None, "source_name": source_name.strip()}

    @retry_transaction
    def create(self, context, request_id, rows, expires_at, meal_type_ids, department_id, event_name=None, source_name=None):
        try:
            identifier = UUID(str(request_id))
        except (TypeError, ValueError):
            raise DomainError("INVALID_REQUEST_ID") from None
        settings = self._settings(expires_at, meal_type_ids, department_id, event_name, source_name)
        clean, errors = validate_rows(rows)
        if errors:
            raise DomainError("VISITOR_ROWS_INVALID")
        canonical = json.dumps(
            {"rows": [{key: row[key] for key in ("sr_no", "person_name", "company_name", "phone", "email")} for row in clean],
             **{key: value.isoformat() if isinstance(value, datetime) else value for key, value in settings.items()
                if key != "source_name"}},
            sort_keys=True, separators=(",", ":"),
        )
        request_hash = hashlib.sha256(canonical.encode()).digest()
        with self.db.transaction() as tx:
            actor = require_actor(tx, context, {"ADMIN"})
            tx.one("SELECT id FROM staff_accounts WHERE id = %s FOR UPDATE", (actor.staff_id,))
            existing = tx.all(
                "SELECT id, request_id, request_hash FROM visitor_qr_batches WHERE created_by_staff_id = %s "
                "AND (request_id = %s OR request_hash = %s) ORDER BY id",
                (actor.staff_id, identifier.bytes, request_hash),
            )
            for row in existing:
                if bytes(row["request_id"]) == identifier.bytes and bytes(row["request_hash"]) != request_hash:
                    raise DomainError("REQUEST_MISMATCH")
            if existing:
                return {"batch_id": existing[0]["id"], "duplicate": True}
            if settings["expires_at"] <= tx.now():
                raise DomainError("INVALID_QR_EXPIRY")
            self.qr.visitor_context(tx, settings["department_id"], settings["meal_type_ids"])
            email_batch_id = None
            if any(row["email"] for row in clean):
                email_batch_id = tx.insert(
                    "INSERT INTO employee_email_batches (request_id, created_by_staff_id, request_hash, created_at) "
                    "VALUES (%s, %s, %s, %s)",
                    (uuid4().bytes, actor.staff_id, request_hash, tx.now()),
                )
            batch_id = tx.insert(
                "INSERT INTO visitor_qr_batches (request_id, request_hash, created_by_staff_id, department_id, event_name, "
                "expires_at, meal_type_ids, row_count, source_name, email_batch_id, created_at) "
                "VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                (identifier.bytes, request_hash, actor.staff_id, settings["department_id"], settings["event_name"],
                 settings["expires_at"], json.dumps(settings["meal_type_ids"]), len(clean), settings["source_name"],
                 email_batch_id, tx.now()),
            )
            aware_expiry = expires_at
            for row in clean:
                details = {"company_name": row["company_name"], "contact_name": row["person_name"], "email": row["email"],
                           "phone": row["phone"], "meal_limit": 1, "meal_type_ids": settings["meal_type_ids"],
                           "department_id": settings["department_id"], "event_name": settings["event_name"]}
                issued = self.qr._issue(tx, actor.staff_id, "MASTER", None, aware_expiry)
                visitor = self.qr._create_master_allocation(tx, actor.staff_id, issued.qr_id, details, batch_id, row["sr_no"])
                self.qr.queue_visitor_email(tx, issued.qr_id, issued.token, details, visitor, aware_expiry,
                                            bulk_batch_id=email_batch_id)
            audit(tx, actor.staff_id, "VISITOR_BATCH_CREATED", "visitor_qr_batches", batch_id,
                  after={"row_count": len(clean), "department_id": settings["department_id"],
                         "event_name": settings["event_name"], "meal_type_ids": settings["meal_type_ids"],
                         "email_batch_id": email_batch_id})
        return {"batch_id": batch_id, "duplicate": False}

    def list(self, context, limit=50, before=None):
        if type(limit) is not int or not 1 <= limit <= 200:
            raise DomainError("INVALID_REPORT_LIMIT")
        cursor = parse_cursor(before)
        condition, values, order = newest_first("b.created_at", "b.id", cursor)
        with self.db.transaction() as tx:
            require_actor(tx, context, {"ADMIN"})
            rows = tx.all(
                "SELECT b.id, b.created_at, b.expires_at, b.event_name, b.row_count, b.source_name, b.meal_type_ids, "
                "d.name AS department_name, s.display_name AS created_by, "
                "(SELECT COUNT(*) FROM master_qr_allocations a WHERE a.batch_id = b.id AND a.exhausted_at IS NOT NULL) AS used_up, "
                "(SELECT COUNT(*) FROM email_queue q WHERE q.bulk_batch_id = b.email_batch_id AND q.status = 'PENDING_APPROVAL') AS emails_pending, "
                "(SELECT COUNT(*) FROM email_queue q WHERE q.bulk_batch_id = b.email_batch_id AND q.status = 'QUEUED') AS emails_queued, "
                "(SELECT COUNT(*) FROM email_queue q JOIN master_qr_allocations a ON a.qr_id = q.qr_id "
                "WHERE a.batch_id = b.id AND q.status = 'SENT') AS emails_sent, "
                "(SELECT COUNT(*) FROM email_queue q JOIN master_qr_allocations a ON a.qr_id = q.qr_id "
                "WHERE a.batch_id = b.id AND q.status = 'FAILED') AS emails_failed "
                "FROM visitor_qr_batches b JOIN departments d ON d.id = b.department_id "
                "JOIN staff_accounts s ON s.id = b.created_by_staff_id"
                + (" WHERE " + condition if condition else "") + order,
                (*values, limit + 1),
            )
            remaining, sent = self.remaining_today(tx)
        items = []
        for row in rows[:limit]:
            item = _record(row)
            item["meal_type_ids"] = json.loads(row["meal_type_ids"]) if isinstance(row["meal_type_ids"], str) else row["meal_type_ids"]
            items.append(item)
        return {"items": items, "next_cursor": next_cursor(rows, limit, "created_at"),
                "email_allowance": {"daily_limit": self.daily_limit, "sent_last_24h": sent, "remaining": remaining}}

    def detail(self, context, batch_id):
        if type(batch_id) is not int or not 1 <= batch_id <= 2**64 - 1:
            raise DomainError("INVALID_VISITOR_BATCH_ID")
        with self.db.transaction() as tx:
            require_actor(tx, context, {"ADMIN"})
            batch = tx.one(
                "SELECT b.id, b.created_at, b.expires_at, b.event_name, b.row_count, b.source_name, b.meal_type_ids, "
                "b.email_batch_id, d.name AS department_name FROM visitor_qr_batches b "
                "JOIN departments d ON d.id = b.department_id WHERE b.id = %s",
                (batch_id,),
            )
            if batch is None:
                raise DomainError("VISITOR_BATCH_NOT_FOUND")
            people = tx.all(
                "SELECT a.qr_id, a.batch_row AS sr_no, a.contact_name AS person_name, a.company_name, a.phone, a.email, "
                "a.exhausted_at, c.revoked_at, c.expires_at, "
                "(SELECT q.id FROM email_queue q WHERE q.qr_id = a.qr_id ORDER BY q.id DESC LIMIT 1) AS email_id, "
                "(SELECT q.status FROM email_queue q WHERE q.qr_id = a.qr_id ORDER BY q.id DESC LIMIT 1) AS email_status, "
                "(SELECT q.last_error FROM email_queue q WHERE q.qr_id = a.qr_id ORDER BY q.id DESC LIMIT 1) AS email_error, "
                "(SELECT COUNT(*) FROM email_queue q WHERE q.qr_id = a.qr_id) AS email_count "
                "FROM master_qr_allocations a JOIN qr_credentials c ON c.id = a.qr_id "
                "WHERE a.batch_id = %s ORDER BY a.batch_row",
                (batch_id,),
            )
            now = tx.now()
        from .email_queue import email_delivery_state

        result = _record(batch)
        result["meal_type_ids"] = json.loads(batch["meal_type_ids"]) if isinstance(batch["meal_type_ids"], str) else batch["meal_type_ids"]
        result.pop("email_batch_id")
        result["people"] = []
        for person in people:
            status = "REVOKED" if person["revoked_at"] else "USED_UP" if person["exhausted_at"] else (
                "EXPIRED" if person["expires_at"] is not None and person["expires_at"] <= now else "ACTIVE")
            email_status, code = (None, None) if person["email_status"] is None else email_delivery_state(
                {"status": person["email_status"], "last_error": person["email_error"]})
            result["people"].append({
                "qr_id": person["qr_id"], "sr_no": person["sr_no"], "person_name": person["person_name"],
                "company_name": person["company_name"], "phone": person["phone"], "email": person["email"],
                "qr_status": status, "email_id": person["email_id"], "email_status": email_status, "email_code": code,
                "resends_left": max(0, RESEND_LIMIT + 1 - person["email_count"]) if person["email"] else 0,
            })
        return result

    def remaining_today(self, tx):
        since = tx.now() - timedelta(days=1)
        sent = self.qr.email_queue.sent_since(tx, since)
        queued = tx.one("SELECT COUNT(*) AS queued FROM email_queue WHERE status = 'QUEUED'")["queued"]
        return max(0, self.daily_limit - sent - queued), sent

    @retry_transaction
    def approve_emails(self, context, batch_id, expected_count):
        if type(batch_id) is not int or not 1 <= batch_id <= 2**64 - 1:
            raise DomainError("INVALID_VISITOR_BATCH_ID")
        if type(expected_count) is not int or not 1 <= expected_count <= MAX_ROWS:
            raise DomainError("EMAIL_COUNT_CONFIRMATION_REQUIRED")
        with self.db.transaction() as tx:
            actor = require_actor(tx, context, {"ADMIN"})
            batch = tx.one("SELECT id, email_batch_id FROM visitor_qr_batches WHERE id = %s", (batch_id,))
            if batch is None:
                raise DomainError("VISITOR_BATCH_NOT_FOUND")
            if batch["email_batch_id"] is None:
                raise DomainError("VISITOR_BATCH_HAS_NO_EMAILS")
            pending = [row["id"] for row in tx.all(
                "SELECT id FROM email_queue WHERE bulk_batch_id = %s AND status = 'PENDING_APPROVAL' ORDER BY id",
                (batch["email_batch_id"],),
            )]
            if not pending:
                raise DomainError("VISITOR_BATCH_NOTHING_TO_SEND")
            if len(pending) != expected_count:
                raise DomainError("EMAIL_COUNT_CHANGED")
            remaining, _ = self.remaining_today(tx)
            if len(pending) > remaining:
                raise DomainError("EMAIL_DAILY_LIMIT_EXCEEDED")
            for employee, credential, queued in self.qr.email_queue._bulk_rows(tx, pending):
                self.qr.email_queue.validated_preview(tx, employee, credential, queued, statuses=("PENDING_APPROVAL",), include_svg=False)
                self.qr.email_queue._approve(tx, actor, queued)
            audit(tx, actor.staff_id, "VISITOR_BATCH_EMAILS_APPROVED", "visitor_qr_batches", batch_id,
                  after={"email_count": len(pending)})
        return {"batch_id": batch_id, "email_ids": pending, "approved_count": len(pending)}

    @retry_transaction
    def prepare_single(self, context, qr_id):
        if type(qr_id) is not int or not 1 <= qr_id <= 2**64 - 1:
            raise DomainError("INVALID_QR_ID")
        with self.db.transaction() as tx:
            actor = require_actor(tx, context, {"ADMIN"})
            credential = tx.one(
                "SELECT id, kind, revoked_at, expires_at, token_ciphertext, token_hash FROM qr_credentials WHERE id = %s FOR UPDATE",
                (qr_id,),
            )
            allocation = tx.one(
                "SELECT qr_id, company_name, contact_name, email, phone, meal_limit, department_id, event_name, exhausted_at "
                "FROM master_qr_allocations WHERE qr_id = %s",
                (qr_id,),
            )
            if credential is None or credential["kind"] != "MASTER" or allocation is None:
                raise DomainError("MASTER_QR_REQUIRED")
            if not allocation["email"]:
                raise DomainError("VISITOR_EMAIL_MISSING")
            if credential["revoked_at"] is not None or allocation["exhausted_at"] is not None or (
                    credential["expires_at"] is not None and credential["expires_at"] <= tx.now()):
                raise DomainError("QR_NOT_ACTIVE")
            latest = tx.one(
                "SELECT id, status, delivery_mode FROM email_queue WHERE qr_id = %s ORDER BY id DESC LIMIT 1 FOR UPDATE",
                (qr_id,),
            )
            if latest is not None and latest["status"] in {"QUEUED", "PENDING_APPROVAL"}:
                raise DomainError("EMAIL_ALREADY_PENDING")
            remaining, _ = self.remaining_today(tx)
            if remaining < 1:
                raise DomainError("EMAIL_DAILY_LIMIT_EXCEEDED")
            if latest is not None and latest["status"] == "DRAFT" and latest["delivery_mode"] == "SINGLE":
                email_id = latest["id"]
            else:
                count = tx.one("SELECT COUNT(*) AS sent FROM email_queue WHERE qr_id = %s AND recipient_email = %s",
                               (qr_id, allocation["email"]))["sent"]
                if count > RESEND_LIMIT:
                    raise DomainError("EMAIL_RESEND_LIMIT_REACHED")
                if credential["token_ciphertext"] is None:
                    raise DomainError("QR_TOKEN_UNAVAILABLE")
                token = self.qr._decrypt_token(credential)
                details = {"email": allocation["email"], "contact_name": allocation["contact_name"],
                           "company_name": allocation["company_name"], "event_name": allocation["event_name"]}
                meal_type_ids = [row["meal_type_id"] for row in tx.all(
                    "SELECT meal_type_id FROM master_qr_meal_types WHERE qr_id = %s ORDER BY meal_type_id", (qr_id,))]
                department = tx.one("SELECT name FROM departments WHERE id = %s", (allocation["department_id"],)) if allocation["department_id"] else None
                names = [row["name"] for row in tx.all(
                    "SELECT name FROM meal_types WHERE id IN (" + ", ".join(["%s"] * len(meal_type_ids)) + ") ORDER BY id",
                    tuple(meal_type_ids))] if meal_type_ids else []
                visitor = {"name": allocation["contact_name"], "company_name": allocation["company_name"],
                           "department_name": department["name"] if department else None,
                           "event_name": allocation["event_name"], "meal_types": names}
                expiry = None if credential["expires_at"] is None else credential["expires_at"].isoformat() + "+00:00"
                email_id = self.qr.email_queue.enqueue_visitor(
                    tx, qr_id, allocation["email"], token, {**visitor, "expires_at": expiry})
                if latest is not None:
                    audit(tx, actor.staff_id, "VISITOR_EMAIL_RESEND_CREATED", "email_queue", email_id,
                          after={"qr_id": qr_id, "previous_email_id": latest["id"]})
            self.qr.email_queue.approve_visitor(tx, actor, email_id, qr_id)
        return email_id
